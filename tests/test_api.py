import csv
import io

from openpyxl import load_workbook

from tests.conftest import RESUMES, resume_bytes

BACKEND_JD = {
    "title": "Backend Engineer",
    "company": "Acme",
    "description": "Requirements\n- 4+ years of experience\n- Bachelor's degree\n- Python and FastAPI or Django\n- PostgreSQL, Docker and AWS\n\nNice to have\n- Kubernetes",
}


def upload(client, names, job_id=None):
    files = [("files", (n, resume_bytes(n), "application/octet-stream")) for n in names]
    return client.post("/api/resumes/upload", files=files, data={"job_id": job_id} if job_id else {})


def test_health(client):
    r = client.get("/api/health")
    assert r.status_code == 200 and r.json()["status"] == "ok"


def test_openapi_docs_are_served(client):
    spec = client.get("/openapi.json").json()
    assert {"/api/jobs", "/api/resumes/upload", "/api/jobs/{job_id}/rankings", "/api/jobs/{job_id}/export"} <= set(spec["paths"])
    assert client.get("/docs").status_code == 200
    assert client.get("/").status_code == 200


def test_create_job_auto_extracts_requirements(client):
    r = client.post("/api/jobs", json=BACKEND_JD)
    assert r.status_code == 201
    job = r.json()
    assert "FastAPI | Django" in job["required_skills"]
    assert job["preferred_skills"] == ["Kubernetes"]
    assert job["min_experience"] == 4.0 and job["education_level"] == "bachelor"


def test_explicit_requirements_override_extraction(client):
    job = client.post("/api/jobs", json={**BACKEND_JD, "required_skills": ["postgres", "Go"], "min_experience": 1, "education_level": "master"}).json()
    assert job["required_skills"] == ["PostgreSQL", "Go"]
    assert job["min_experience"] == 1 and job["education_level"] == "master"


def test_job_validation_errors(client):
    assert client.post("/api/jobs", json={"title": "x", "description": "short"}).status_code == 422
    assert client.post("/api/jobs", json={**BACKEND_JD, "education_level": "wizard"}).status_code == 422


def test_job_crud(client):
    jid = client.post("/api/jobs", json=BACKEND_JD).json()["id"]
    assert client.get(f"/api/jobs/{jid}").json()["title"] == "Backend Engineer"
    r = client.put(f"/api/jobs/{jid}", json={"title": "Staff Backend Engineer", "required_skills": ["Python"]})
    assert r.json()["title"] == "Staff Backend Engineer" and r.json()["required_skills"] == ["Python"]
    assert client.delete(f"/api/jobs/{jid}").status_code == 204
    r = client.get(f"/api/jobs/{jid}")
    assert r.status_code == 404 and r.json()["error"]["code"] == "not_found"


def test_parse_description_endpoint(client):
    r = client.post("/api/jobs/parse-description", json={"description": BACKEND_JD["description"]})
    assert r.status_code == 200 and "Python" in r.json()["required_skills"]


def test_upload_ranks_candidates_and_returns_scores(client):
    job_id = client.post("/api/jobs", json=BACKEND_JD).json()["id"]
    r = upload(client, ["aarav_mehta.pdf", "vikram_singh.docx", "sofia_alvarez.pdf"], job_id)
    body = r.json()
    assert (body["processed"], body["duplicates"], body["failed"]) == (3, 0, 0)
    scores = {i["name"]: i["score"] for i in body["results"]}
    assert scores["Aarav Mehta"] > scores["Sofia Alvarez"] > scores["Vikram Singh"]

    rk = client.get(f"/api/jobs/{job_id}/rankings").json()
    assert [i["candidate"]["name"] for i in rk["items"]] == ["Aarav Mehta", "Sofia Alvarez", "Vikram Singh"]
    assert [i["rank"] for i in rk["items"]] == [1, 2, 3]
    assert all(0 <= i["score"] <= 100 for i in rk["items"])
    top = rk["items"][0]
    assert top["skill_analysis"]["missing_required"] == []
    assert set(top["breakdown"]) == {"skills", "semantic", "experience", "education"}


def test_new_job_scores_existing_candidates_and_new_resume_scores_existing_jobs(client):
    upload(client, ["aarav_mehta.pdf"])
    jid = client.post("/api/jobs", json=BACKEND_JD).json()["id"]
    assert client.get(f"/api/jobs/{jid}/rankings").json()["total"] == 1
    upload(client, ["sofia_alvarez.pdf"])
    assert client.get(f"/api/jobs/{jid}/rankings").json()["total"] == 2


def test_duplicate_upload_is_detected(client):
    upload(client, ["aarav_mehta.pdf"])
    r = upload(client, ["aarav_mehta.pdf"]).json()
    assert r["duplicates"] == 1 and r["processed"] == 0
    assert len(client.get("/api/candidates").json()) == 1


def test_same_person_new_file_updates_instead_of_duplicating(client):
    upload(client, ["aarav_mehta.pdf"])
    from app.services.parser import parse_resume
    from docx import Document

    d = Document()
    for line in ["Aarav Mehta", "aarav.mehta@example.com", "Skills", "Python, Rust, PostgreSQL", "Experience", "Engineer 2015 - 2025 at Acme"]:
        d.add_paragraph(line)
    buf = io.BytesIO()
    d.save(buf)
    r = client.post("/api/resumes/upload", files=[("files", ("aarav_new.docx", buf.getvalue(), "application/octet-stream"))]).json()
    assert r["results"][0]["status"] == "updated"
    cands = client.get("/api/candidates").json()
    assert len(cands) == 1 and "Rust" in cands[0]["skills"]


def test_bad_files_fail_individually_without_aborting_the_batch(client):
    files = [
        ("files", ("notes.txt", b"hello", "text/plain")),
        ("files", ("fake.pdf", b"not a pdf at all", "application/pdf")),
        ("files", ("aarav_mehta.pdf", resume_bytes("aarav_mehta.pdf"), "application/pdf")),
    ]
    r = client.post("/api/resumes/upload", files=files)
    assert r.status_code == 200
    body = r.json()
    assert (body["processed"], body["failed"]) == (1, 2)
    errs = {i["filename"]: i["error"] for i in body["results"] if i["error"]}
    assert "only PDF and DOCX" in errs["notes.txt"] and "not a valid PDF" in errs["fake.pdf"]


def test_upload_with_unknown_job_id_is_404(client):
    assert upload(client, ["aarav_mehta.pdf"], job_id=999).status_code == 404


def test_scanned_resume_is_processed_via_ocr(client):
    import shutil

    import pytest

    if not shutil.which("tesseract"):
        pytest.skip("Tesseract not installed")
    r = upload(client, ["nikhil_verma_scanned.pdf"]).json()
    assert r["processed"] == 1 and r["results"][0]["name"] == "Nikhil Verma"
    assert r["results"][0]["warnings"]


def test_filters(seeded):
    jid = 1
    total = seeded.get(f"/api/jobs/{jid}/rankings").json()["total"]
    assert total == 15

    def names(**params):
        return {i["candidate"]["name"] for i in seeded.get(f"/api/jobs/{jid}/rankings", params=params).json()["items"]}

    assert names(skills="kubernetes,terraform") == {"Aarav Mehta", "Carlos Rivera"}
    assert "Aarav Mehta" in names(skills="kubernetes,react", skill_mode="any") and "Meera Krishnan" in names(skills="kubernetes,react", skill_mode="any")
    assert {"Aarav Mehta", "Sofia Alvarez"} <= names(skills="postgres", min_experience=8)
    assert "Liam Chen" not in names(skills="postgres", min_experience=8)
    assert all(n in names(min_experience=8) for n in ("Aarav Mehta", "Carlos Rivera"))
    assert "Liam Chen" not in names(min_experience=3)
    assert names(education="master") >= {"Aarav Mehta", "Ananya Iyer", "Vikram Singh"}
    assert "Sofia Alvarez" not in names(education="master")
    assert names(q="okafor") == {"Daniel Okafor"}
    assert all(s >= 70 for s in [i["score"] for i in seeded.get(f"/api/jobs/{jid}/rankings", params={"min_score": 70}).json()["items"]])


def test_ranks_stay_global_when_filtering(seeded):
    full = seeded.get("/api/jobs/1/rankings").json()["items"]
    expected = next(i["rank"] for i in full if i["candidate"]["name"] == "Sofia Alvarez")
    items = seeded.get("/api/jobs/1/rankings", params={"q": "sofia"}).json()["items"]
    assert len(items) == 1 and items[0]["rank"] == expected > 1


def test_invalid_filter_values_are_rejected(seeded):
    assert seeded.get("/api/jobs/1/rankings", params={"education": "wizard"}).status_code == 422
    assert seeded.get("/api/jobs/1/rankings", params={"min_score": 500}).status_code == 422


def test_shortlist_status_flow_and_export(seeded):
    items = seeded.get("/api/jobs/1/rankings").json()["items"]
    top, second = items[0]["match_id"], items[1]["match_id"]
    assert seeded.patch(f"/api/matches/{top}", json={"status": "shortlisted"}).json()["status"] == "shortlisted"
    seeded.patch(f"/api/matches/{second}", json={"status": "rejected"})
    assert seeded.patch(f"/api/matches/{top}", json={"status": "bogus"}).status_code == 422

    assert seeded.get("/api/jobs/1").json()["shortlisted_count"] == 1
    sl = seeded.get("/api/jobs/1/rankings", params={"status": "shortlisted"}).json()
    assert sl["total"] == 1

    r = seeded.get("/api/jobs/1/export", params={"format": "csv", "shortlisted_only": True})
    assert r.status_code == 200 and "attachment" in r.headers["content-disposition"]
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert rows[0][:5] == ["Rank", "Name", "Email", "Phone", "Match score (%)"]
    assert len(rows) == 2 and rows[1][1] == items[0]["candidate"]["name"]


def test_excel_export_is_a_valid_styled_workbook(seeded):
    r = seeded.get("/api/jobs/1/export", params={"format": "xlsx"})
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Shortlist", "Job requirements"]
    ws = wb["Shortlist"]
    assert ws["A1"].value.startswith("Senior Backend Engineer")
    assert ws.max_row == 4 + 15
    assert ws["B5"].value == "Aarav Mehta" and ws.freeze_panes == "C5"
    assert wb["Job requirements"]["A3"].value == "Required skills"


def test_export_respects_filters(seeded):
    r = seeded.get("/api/jobs/1/export", params={"format": "csv", "skills": "kubernetes", "min_score": 50})
    rows = list(csv.reader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert 1 < len(rows) < 10


def test_match_detail_candidate_and_resume_download(seeded):
    m = seeded.get("/api/jobs/1/rankings").json()["items"][0]
    d = seeded.get(f"/api/matches/{m['match_id']}").json()
    c = d["candidate"]
    assert c["education"] and c["experience"] and c["certifications"] and c["parse_method"]
    file = seeded.get(f"/api/candidates/{c['id']}/resume")
    assert file.status_code == 200 and file.content[:4] == b"%PDF"
    assert seeded.get("/api/candidates/9999").status_code == 404


def test_delete_candidate_removes_matches(seeded):
    cid = seeded.get("/api/jobs/1/rankings").json()["items"][0]["candidate"]["id"]
    assert seeded.delete(f"/api/candidates/{cid}").status_code == 204
    assert seeded.get("/api/jobs/1/rankings").json()["total"] == 14


def test_interview_questions_are_tailored(seeded):
    ranking = seeded.get("/api/jobs/1/rankings").json()["items"]
    sofia = next(i for i in ranking if i["candidate"]["name"] == "Sofia Alvarez")
    r = seeded.get(f"/api/matches/{sofia['match_id']}/interview-questions", params={"count": 8}).json()
    assert r["source"] == "template" and 3 <= len(r["questions"]) <= 8
    cats = {q["category"] for q in r["questions"]}
    assert {"Technical depth", "Gap probing", "Behavioural"} <= cats
    assert any("Microservices" in q["question"] for q in r["questions"] if q["category"] == "Gap probing")
    assert all(q["question"] and q["rationale"] for q in r["questions"])


def test_notify_is_a_dry_run_without_smtp_and_marks_the_candidate(seeded):
    mid = seeded.get("/api/jobs/1/rankings").json()["items"][0]["match_id"]
    r = seeded.post(f"/api/matches/{mid}/notify", json={}).json()
    assert r["status"] == "dry_run" and r["to"] == "aarav.mehta@example.com"
    assert seeded.get(f"/api/matches/{mid}").json()["notified_at"] is not None


def test_bulk_notify_only_emails_shortlisted_and_only_once(seeded):
    items = seeded.get("/api/jobs/1/rankings").json()["items"]
    for it in items[:2]:
        seeded.patch(f"/api/matches/{it['match_id']}", json={"status": "shortlisted"})
    assert len(seeded.post("/api/jobs/1/notify-shortlisted", json={}).json()["results"]) == 2
    assert seeded.post("/api/jobs/1/notify-shortlisted", json={}).json()["results"] == []


def test_notify_with_custom_template(seeded, monkeypatch):
    sent = {}
    from app.services import notifier

    def fake_send(to, subject, body):
        sent.update(to=to, subject=subject, body=body)
        return "sent"

    monkeypatch.setattr("app.api.routes.notifier.send_email", fake_send)
    mid = seeded.get("/api/jobs/1/rankings").json()["items"][0]["match_id"]
    r = seeded.post(f"/api/matches/{mid}/notify", json={"subject": "Hi {name}", "body": "Role: {job_title}{company_part}"}).json()
    assert r["status"] == "sent" and sent["subject"] == "Hi Aarav" and sent["body"] == "Role: Senior Backend Engineer at Northwind Labs"
    assert notifier.DEFAULT_SUBJECT


def test_smtp_failure_is_reported_not_raised(seeded, monkeypatch):
    from app.exceptions import EmailDeliveryError

    def boom(*a, **k):
        raise EmailDeliveryError("SMTP down")

    monkeypatch.setattr("app.api.routes.notifier.send_email", boom)
    mid = seeded.get("/api/jobs/1/rankings").json()["items"][0]["match_id"]
    r = seeded.post(f"/api/matches/{mid}/notify", json={})
    assert r.status_code == 200 and r.json()["status"] == "failed" and "SMTP down" in r.json()["detail"]


def test_semantic_search_finds_the_right_people(seeded):
    def top(q):
        return [h["candidate"]["name"] for h in seeded.get("/api/search", params={"q": q, "job_id": 1}).json()["hits"]]

    assert top("built LLM apps with retrieval augmented generation")[0] == "Ananya Iyer"
    assert top("kubernetes and terraform infrastructure automation")[:2] == ["Carlos Rivera", "Aarav Mehta"]
    assert top("payroll and recruitment")[0] == "Vikram Singh"
    assert top("quantum chemistry laboratory") == []
    hit = seeded.get("/api/search", params={"q": "react accessibility design systems", "job_id": 3}).json()["hits"][0]
    assert hit["candidate"]["name"] == "Meera Krishnan" and hit["job_score"] > 80


def test_top_skills_endpoint(seeded):
    skills = seeded.get("/api/skills/top", params={"limit": 5}).json()
    assert len(skills) == 5 and skills[0]["count"] >= skills[-1]["count"]


def test_seed_is_idempotent(client):
    first = client.post("/api/demo/seed").json()
    second = client.post("/api/demo/seed").json()
    assert first == {"jobs_created": 3, "resumes_created": 15} and second == {"jobs_created": 0, "resumes_created": 0}


def test_all_sample_resumes_are_processed(client):
    names = sorted(p.name for p in RESUMES.iterdir())
    body = upload(client, names).json()
    assert body["failed"] == 0 and body["processed"] == len(names)

