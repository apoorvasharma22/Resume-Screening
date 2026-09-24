from __future__ import annotations

import csv
import html
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, SimpleDocTemplate, Spacer, FrameBreak

OUT = Path(__file__).resolve().parent.parent / "sample_data"

R = []


def resume(**kw):
    R.append(kw)


resume(
    file="aarav_mehta.pdf", layout="classic_pdf", name="Aarav Mehta", email="aarav.mehta@example.com", phone="+91 98765 43210",
    location="Bengaluru, India", title="Senior Backend Engineer",
    summary="Backend engineer with 9+ years of experience building high-throughput APIs and event-driven systems in Python. Led a team of six and cut p95 latency by 60% across the payments platform.",
    skills={"Languages": ["Python", "SQL", "Bash", "Go"], "Frameworks": ["FastAPI", "Django", "Celery"], "Data": ["PostgreSQL", "Redis", "Kafka", "Elasticsearch"],
            "Cloud & DevOps": ["AWS", "Docker", "Kubernetes", "Terraform", "GitHub Actions", "CI/CD"]},
    jobs=[("Staff Software Engineer", "PayNest Technologies", "Mar 2021", "Present",
           ["Designed a microservices architecture in Python and FastAPI processing 4M payment events per day", "Introduced Kafka-based event streaming and PostgreSQL partitioning, reducing p95 latency by 60%",
            "Mentored 5 engineers and ran code reviews; owned CI/CD pipelines on GitHub Actions and Kubernetes"]),
          ("Senior Backend Developer", "CloudKart", "Jan 2019", "Feb 2021",
           ["Built REST APIs with Django and Redis caching serving 2M monthly users", "Migrated services to Docker and AWS ECS with Terraform"]),
          ("Software Engineer", "InfoStack Solutions", "Aug 2017", "Dec 2018", ["Developed internal tooling in Python and PostgreSQL"])],
    edu=[("M.Tech in Computer Science", "Indian Institute of Technology Delhi", "2015 - 2017"), ("B.Tech in Information Technology", "Delhi Technological University", "2011 - 2015")],
    certs=["AWS Certified Solutions Architect - Associate (2022)", "Certified Kubernetes Administrator (CKA) 2023"],
)
resume(
    file="sofia_alvarez.pdf", layout="sidebar_pdf", name="Sofia Alvarez", email="sofia.alvarez@example.com", phone="+34 612 345 678",
    location="Madrid, Spain", title="Backend Engineer",
    summary="Backend engineer with 8 years of experience building web platforms with Python and Django. Comfortable owning features from database design to production deployment on AWS.",
    skills={"Languages": ["Python", "SQL", "JavaScript"], "Frameworks": ["Django", "Django REST Framework", "Celery"], "Data": ["PostgreSQL", "Redis", "MongoDB"],
            "Tools": ["Docker", "AWS", "Jenkins", "Git", "Linux"]},
    jobs=[("Backend Engineer", "Mercado Digital", "Feb 2021", "Present", ["Built REST APIs with Django REST Framework and PostgreSQL for a marketplace with 500k users", "Containerised services with Docker and deployed to AWS EC2 through Jenkins pipelines", "Added Redis caching, cutting response times by 40%"]),
          ("Python Developer", "Andes Software", "Sep 2019", "Jan 2021", ["Developed Celery workers and REST APIs for a logistics tracking product", "Wrote pytest suites reaching 85% coverage"]),
          ("Junior Developer", "WebForma", "Jun 2018", "Aug 2019", ["Maintained Django sites and MySQL databases"])],
    edu=[("B.Sc. in Computer Science", "Universidad Politécnica de Madrid", "2014 - 2018")], certs=["AWS Certified Cloud Practitioner"],
)
resume(
    file="daniel_okafor.docx", layout="docx_classic", name="Daniel Okafor", email="daniel.okafor@example.com", phone="+44 7700 900123",
    location="London, UK", title="Software Engineer",
    summary="Software engineer with 8 years of experience in Java and Spring Boot microservices. Some Python scripting for automation and data tasks.",
    skills={"Languages": ["Java", "Kotlin", "SQL", "Python (basic scripting)"], "Frameworks": ["Spring Boot", "Hibernate"], "Data": ["MySQL", "Kafka", "RabbitMQ"], "Cloud & DevOps": ["AWS", "Docker", "Jenkins", "CI/CD"]},
    jobs=[("Software Engineer", "Thames Financial Systems", "May 2020", "Present", ["Built Spring Boot microservices and REST APIs handling trade settlement", "Integrated Kafka topics and MySQL persistence; deployed with Docker on AWS", "Wrote Python scripts to automate reporting"]),
          ("Junior Java Developer", "Northbridge Consulting", "Jul 2018", "Apr 2020", ["Developed Hibernate data-access layers and JUnit tests"])],
    edu=[("BSc in Software Engineering", "University of Manchester", "2014 - 2018")], certs=["Oracle Certified Professional, Java SE 11"],
)
resume(
    file="priya_nair.docx", layout="docx_table", name="Priya Nair", email="priya.nair@example.com", phone="+91 91234 56780",
    location="Pune, India", title="Full Stack Developer",
    summary="Full stack developer with 6 years of experience building products with Flask and React. Enjoys clean APIs and readable code.",
    skills={"Languages": ["Python", "JavaScript", "SQL"], "Frameworks": ["Flask", "React"], "Data": ["PostgreSQL"], "Tools": ["Docker", "Git", "REST APIs"]},
    jobs=[("Full Stack Developer", "BrightLeaf Labs", "Jun 2022", "Present", ["Built Flask REST APIs backed by PostgreSQL and a React front end", "Dockerised the stack and set up basic GitHub CI"]),
          ("Software Developer", "Punetech Services", "Aug 2020", "May 2022", ["Developed Python scripts and Flask endpoints for an HR portal"])],
    edu=[("B.Tech in Computer Engineering", "Savitribai Phule Pune University", "2016 - 2020")], certs=[],
)
resume(
    file="liam_chen.pdf", layout="compact_pdf", name="Liam Chen", email="liam.chen@example.com", phone="+1 415 555 0142",
    location="San Francisco, CA", title="Junior Backend Developer",
    summary="Junior developer with about 1.5 years of professional experience in Python web development.",
    skills={"Languages": ["Python", "SQL"], "Frameworks": ["Flask"], "Data": ["SQLite"], "Tools": ["Git"]},
    jobs=[("Junior Backend Developer", "Startlane Inc.", "03/2025", "Present", ["Built Flask endpoints and SQLite-backed features", "Fixed bugs and wrote unit tests with pytest"])],
    edu=[("B.S. in Computer Science", "San Jose State University", "2020 - 2024")], certs=[],
)
resume(
    file="ananya_iyer.pdf", layout="sidebar_pdf", name="Ananya Iyer", email="ananya.iyer@example.com", phone="+91 99000 11223",
    location="Hyderabad, India", title="Machine Learning Engineer",
    summary="Machine learning engineer with 6 years of experience shipping NLP and LLM applications to production. Strong in PyTorch, retrieval-augmented generation and MLOps.",
    skills={"Languages": ["Python", "SQL"], "ML": ["PyTorch", "scikit-learn", "NLP", "Machine Learning", "Deep Learning", "Hugging Face"], "LLM": ["LLMs", "RAG", "LangChain", "Prompt Engineering", "FAISS"],
            "MLOps & Cloud": ["MLflow", "Docker", "AWS", "Airflow", "Spark"]},
    jobs=[("Machine Learning Engineer", "Lexica AI", "Apr 2022", "Present", ["Built a RAG assistant with LangChain, FAISS and LLMs used by 20k customers", "Fine-tuned transformer models with PyTorch and Hugging Face, improving F1 by 12 points", "Deployed models with Docker on AWS and tracked experiments in MLflow"]),
          ("Data Scientist", "Retailytics", "Jul 2020", "Mar 2022", ["Developed NLP classifiers with scikit-learn and PyTorch", "Built Spark and Airflow data pipelines"])],
    edu=[("M.S. in Machine Learning", "International Institute of Information Technology Hyderabad", "2018 - 2020"), ("B.Tech in Computer Science", "NIT Warangal", "2014 - 2018")],
    certs=["TensorFlow Developer Certificate", "AWS Certified Machine Learning - Specialty"],
)
resume(
    file="rohan_kapoor.pdf", layout="classic_pdf", name="Rohan Kapoor", email="rohan.kapoor@example.com", phone="+91 98111 22334",
    location="Gurugram, India", title="Data Scientist",
    summary="Data scientist with 5 years of experience in predictive modelling, experimentation and dashboarding for e-commerce and telecom.",
    skills={"Languages": ["Python", "SQL", "R"], "Libraries": ["Pandas", "NumPy", "scikit-learn", "TensorFlow", "Matplotlib"], "Tools": ["Tableau", "Spark", "Jupyter", "Git"], "Methods": ["Statistics", "A/B Testing", "Feature Engineering", "Time Series"]},
    jobs=[("Data Scientist", "Telenova Communications", "Jan 2023", "Present", ["Built churn prediction models with scikit-learn and TensorFlow, lifting retention by 4%", "Designed A/B tests and Tableau dashboards for the growth team"]),
          ("Associate Data Scientist", "ShopSphere", "Jul 2021", "Dec 2022", ["Forecasted demand using time series models in Python and Spark", "Wrote SQL pipelines and feature engineering jobs"])],
    edu=[("M.Sc. in Statistics", "University of Delhi", "2019 - 2021"), ("B.Sc. (Hons) Mathematics", "St. Stephen's College", "2016 - 2019")], certs=["Google Data Analytics Certificate"],
)
resume(
    file="emma_johansson.docx", layout="docx_classic", name="Emma Johansson", email="emma.johansson@example.com", phone="+46 70 123 45 67",
    location="Stockholm, Sweden", title="Data Analyst",
    summary="Data analyst with 5 years of experience turning business questions into SQL analyses and Power BI dashboards. Basic Python for data cleaning.",
    skills={"Analysis": ["SQL", "Excel", "Data Analysis", "Data Visualization"], "BI": ["Power BI", "Tableau"], "Programming": ["Python (pandas basics)"]},
    jobs=[("Data Analyst", "Nordic Retail Group", "Sep 2023", "Present", ["Created Power BI dashboards and SQL reports for 120 stores", "Automated Excel reporting saving 15 hours a week"]),
          ("Junior Analyst", "Fjord Insights", "Aug 2021", "Aug 2023", ["Cleaned datasets with Python and pandas; produced weekly KPI decks"])],
    edu=[("B.A. in Business Analytics", "Stockholm University", "2018 - 2021")], certs=["Microsoft Certified: Power BI Data Analyst Associate"],
)
resume(
    file="meera_krishnan.pdf", layout="compact_pdf", name="Meera Krishnan", email="meera.krishnan@example.com", phone="+91 97777 88990",
    location="Chennai, India", title="Senior Frontend Engineer",
    summary="Frontend engineer with 10 years of experience building accessible, high-performance React and TypeScript applications for e-commerce and fintech.",
    skills={"Languages": ["TypeScript", "JavaScript", "HTML5", "CSS3"], "Frameworks": ["React", "Next.js", "Redux", "GraphQL"], "Testing": ["Jest", "Cypress"], "Tooling": ["Tailwind CSS", "Storybook", "Webpack", "Git", "CI/CD"], "Practices": ["Accessibility", "Responsive Design", "Web Performance"]},
    jobs=[("Senior Frontend Engineer", "Kite Commerce", "02/2021", "Present", ["Led migration to Next.js and TypeScript, improving Core Web Vitals by 35%", "Built a Storybook design system and Tailwind CSS component library", "Championed accessibility (WCAG) and Jest/Cypress test coverage"]),
          ("Frontend Developer", "Zencart", "06/2018", "01/2021", ["Developed React and Redux storefronts consuming GraphQL and REST APIs"]),
          ("Web Developer", "PixelCraft Studio", "07/2016", "05/2018", ["Built responsive sites with HTML, CSS, JavaScript and jQuery"])],
    edu=[("B.Tech in Information Technology", "Anna University", "2012 - 2016")], certs=[],
)
resume(
    file="jonas_weber.docx", layout="docx_table", name="Jonas Weber", email="jonas.weber@example.com", phone="+49 151 2345 6789",
    location="Berlin, Germany", title="Frontend Developer",
    summary="Frontend developer with 5 years of experience building enterprise dashboards with Angular and TypeScript.",
    skills={"Languages": ["TypeScript", "JavaScript", "HTML", "CSS", "Sass"], "Frameworks": ["Angular", "RxJS"], "Testing": ["Jest", "Jasmine"], "Tools": ["Git", "Webpack"]},
    jobs=[("Frontend Developer", "Spree Analytics", "Mar 2023", "Present", ["Built Angular dashboards in TypeScript consuming REST APIs", "Unit tested components with Jest"]),
          ("Junior Web Developer", "Berlinweb GmbH", "Oct 2021", "Feb 2023", ["Implemented responsive layouts with HTML, Sass and JavaScript"])],
    edu=[("B.Sc. in Media Informatics", "Technische Universität Berlin", "2018 - 2021")], certs=[],
)
resume(
    file="zara_ahmed.pdf", layout="classic_pdf", name="Zara Ahmed", email="zara.ahmed@example.com", phone="+971 50 123 4567",
    location="Dubai, UAE", title="UI Developer",
    summary="UI developer with 6 years of experience turning designs into responsive web pages using React and modern CSS.",
    skills={"Languages": ["JavaScript", "HTML", "CSS"], "Frameworks": ["React", "jQuery", "Bootstrap"], "Tools": ["Git", "Figma", "REST APIs"]},
    jobs=[("UI Developer", "Gulf Digital Agency", "Jan 2022", "Present", ["Built React components and Bootstrap layouts for client sites", "Integrated REST APIs and fixed cross-browser issues"]),
          ("Web Designer", "Oasis Media", "Feb 2020", "Dec 2021", ["Produced responsive HTML/CSS pages with jQuery from Figma designs"])],
    edu=[("B.Sc. in Information Systems", "American University of Sharjah", "2015 - 2019")], certs=[],
)
resume(
    file="vikram_singh.docx", layout="docx_classic", name="Vikram Singh", email="vikram.singh@example.com", phone="+91 98220 33445",
    location="Mumbai, India", title="HR Manager",
    summary="HR manager with 10 years of experience in recruitment, payroll and employee engagement for mid-size technology companies.",
    skills={"HR": ["Recruitment", "Payroll", "Onboarding", "Employee Engagement", "HRIS"], "Tools": ["Excel", "SAP"], "Soft skills": ["Leadership", "Communication", "Stakeholder Management"]},
    jobs=[("HR Manager", "Vertex Solutions", "Apr 2019", "Present", ["Led recruitment for 120 hires a year and managed payroll for 600 employees", "Rolled out an HRIS and onboarding programme"]),
          ("HR Executive", "Sundaram Group", "Jun 2016", "Mar 2019", ["Coordinated talent acquisition and employee engagement events"])],
    edu=[("MBA in Human Resource Management", "Symbiosis Institute of Business Management", "2014 - 2016")], certs=["SHRM-CP"],
)
resume(
    file="hannah_lee.pdf", layout="sidebar_pdf", name="Hannah Lee", email="hannah.lee@example.com", phone="+65 8123 4567",
    location="Singapore", title="Software Engineering Graduate",
    summary="Recent computer science graduate with a six-month backend internship and a strong interest in Python web development.",
    skills={"Languages": ["Python", "Java", "SQL"], "Frameworks": ["Django"], "Tools": ["Git", "Docker (basics)"], "Concepts": ["Data Structures", "OOP", "REST APIs"]},
    jobs=[("Backend Intern", "Merlion Payments", "Jan 2025", "Jun 2025", ["Built Django REST endpoints and PostgreSQL migrations for a payments dashboard", "Wrote unit tests and fixed 30+ bugs"])],
    edu=[("B.Tech in Computer Science", "National University of Singapore", "2021 - 2025")], certs=[],
)
resume(
    file="carlos_rivera.pdf", layout="compact_pdf", name="Carlos Rivera", email="carlos.rivera@example.com", phone="+52 55 1234 5678",
    location="Mexico City, Mexico", title="Senior DevOps Engineer",
    summary="DevOps engineer with 11 years of experience automating cloud infrastructure and delivery pipelines on AWS and Kubernetes. Writes Python and Bash tooling.",
    skills={"Cloud": ["AWS", "Terraform", "CloudFormation"], "Containers": ["Docker", "Kubernetes", "Helm"], "CI/CD": ["Jenkins", "GitHub Actions", "GitLab CI"], "Observability": ["Prometheus", "Grafana", "ELK Stack"], "Scripting": ["Python", "Bash", "Linux"]},
    jobs=[("Senior DevOps Engineer", "Azteca Cloud", "05/2020", "Present", ["Managed 40 microservices on Kubernetes and AWS EKS using Terraform and Helm", "Built CI/CD pipelines in Jenkins and GitHub Actions, cutting release time from 2 days to 2 hours", "Set up Prometheus and Grafana monitoring"]),
          ("DevOps Engineer", "Maya Systems", "03/2017", "04/2020", ["Automated deployments with Ansible and Docker; wrote Python and Bash tooling"]),
          ("Systems Administrator", "Telmex Servicios", "01/2015", "02/2017", ["Administered Linux servers and networking"])],
    edu=[("B.E. in Computer Systems Engineering", "Instituto Tecnológico de Monterrey", "2010 - 2014")], certs=["Certified Kubernetes Administrator (CKA)", "HashiCorp Certified: Terraform Associate"],
)

resume(
    file="nikhil_verma_scanned.pdf", layout="scanned_pdf", name="Nikhil Verma", email="nikhil.verma@example.com", phone="+91 98989 76767",
    location="Noida, India", title="Python Backend Developer",
    summary="Backend developer with 7 years of experience building Django and FastAPI services on AWS. This resume is an image-only scan, used to demonstrate OCR.",
    skills={"Languages": ["Python", "SQL"], "Frameworks": ["Django", "FastAPI"], "Data": ["PostgreSQL", "Redis"], "Cloud & DevOps": ["AWS", "Docker", "CI/CD", "Git"]},
    jobs=[("Backend Developer", "Orbit Commerce", "Feb 2022", "Present", ["Built REST APIs in FastAPI and Django backed by PostgreSQL", "Containerised services with Docker and deployed to AWS with CI/CD pipelines"]),
          ("Software Developer", "Techvista", "Jul 2019", "Jan 2022", ["Developed Django applications and Redis caching layers"])],
    edu=[("B.Tech in Computer Science", "Jaypee Institute of Information Technology", "2015 - 2019")], certs=["AWS Certified Developer - Associate"],
)

JOBS = {
    "01_senior_backend_engineer.txt": """Title: Senior Backend Engineer
Company: Northwind Labs
Location: Remote (India / EU)

About the role
We are building the payments platform behind thousands of online stores and need a senior backend engineer to design and scale our core services. You will own services end to end, from database design to production operations, and mentor other engineers.

Responsibilities
- Design and build scalable REST APIs and microservices in Python using FastAPI or Django
- Model data in PostgreSQL and optimise performance
- Containerise and deploy services with Docker on AWS
- Maintain CI/CD pipelines and improve reliability

Requirements
- 5+ years of professional software engineering experience
- Bachelor's degree in Computer Science or equivalent
- Strong Python and experience with REST APIs and microservices
- Solid PostgreSQL and SQL skills
- Hands-on experience with Docker, AWS and CI/CD

Nice to have
- Kubernetes and Terraform
- Redis and Kafka experience
- Leadership or mentoring experience
""",
    "02_machine_learning_engineer.txt": """Title: Machine Learning Engineer
Company: Helix Intelligence
Location: Hyderabad, India (hybrid)

About the role
Join our applied AI team building language products used by enterprise customers. You will take models from research to reliable production services.

Responsibilities
- Build and deploy NLP and LLM-powered features
- Train, evaluate and monitor machine learning models
- Work with data engineers to build feature pipelines

Requirements
- 3+ years of experience in machine learning or data science
- Master's degree in Computer Science, Statistics or a related field
- Strong Python, scikit-learn and PyTorch or TensorFlow
- Experience with NLP and SQL
- Docker for packaging and deploying models

Preferred
- MLOps and MLflow experience
- Experience with LLMs, RAG and LangChain
- AWS and Spark experience
""",
    "03_senior_frontend_engineer.txt": """Title: Senior Frontend Engineer
Company: Lumen Commerce
Location: Remote

About the role
We are rebuilding our storefront experience and need a senior frontend engineer who cares about performance, accessibility and craft.

Responsibilities
- Build and maintain React and TypeScript applications
- Consume REST APIs and GraphQL services
- Write reliable automated tests and improve developer tooling

Requirements
- 4+ years of experience building production web applications
- Bachelor's degree or equivalent experience
- Expert React, TypeScript and JavaScript
- Strong HTML and CSS
- Experience with REST APIs, Git and Jest

Nice to have
- Next.js and GraphQL
- Tailwind CSS and Storybook
- Accessibility and CI/CD experience
""",
}


INK, MUTED, ACCENT = colors.HexColor("#1B2733"), colors.HexColor("#5B6B75"), colors.HexColor("#0F6E73")


def esc(s):
    return html.escape(s, quote=False)


def styles():
    return {
        "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=INK, alignment=1),
        "name_left": ParagraphStyle("nl", fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=INK),
        "title": ParagraphStyle("title", fontName="Helvetica", fontSize=11.5, leading=15, textColor=ACCENT, alignment=1),
        "title_left": ParagraphStyle("tl", fontName="Helvetica", fontSize=11.5, leading=15, textColor=ACCENT),
        "contact": ParagraphStyle("c", fontName="Helvetica", fontSize=9, leading=12, textColor=MUTED, alignment=1),
        "contact_left": ParagraphStyle("cl", fontName="Helvetica", fontSize=9, leading=12, textColor=MUTED),
        "h": ParagraphStyle("h", fontName="Helvetica-Bold", fontSize=10.5, leading=14, textColor=ACCENT, spaceBefore=9, spaceAfter=3),
        "body": ParagraphStyle("b", fontName="Helvetica", fontSize=9.5, leading=13, textColor=INK),
        "role": ParagraphStyle("r", fontName="Helvetica-Bold", fontSize=10, leading=13, textColor=INK, spaceBefore=4),
        "bullet": ParagraphStyle("bl", fontName="Helvetica", fontSize=9.3, leading=12.5, textColor=INK, leftIndent=10, bulletIndent=0),
        "small": ParagraphStyle("s", fontName="Helvetica", fontSize=8.8, leading=12, textColor=INK),
    }


def _contact(r):
    return f"{r['email']}  |  {r['phone']}  |  {r['location']}"


def build_classic_pdf(r, path):
    s = styles()
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=14 * mm, title=f"{r['name']} - Resume")
    f = [Paragraph(esc(r["name"]), s["name_left"]), Paragraph(esc(r["title"]), s["title_left"]), Paragraph(esc(_contact(r)), s["contact_left"]),
         Paragraph("PROFESSIONAL SUMMARY", s["h"]), Paragraph(esc(r["summary"]), s["body"]), Paragraph("EXPERIENCE", s["h"])]
    for t, c, a, b, bl in r["jobs"]:
        f.append(Paragraph(f"{esc(t)}, {esc(c)}  |  {a} - {b}", s["role"]))
        f += [Paragraph(esc(x), s["bullet"], bulletText="•") for x in bl]
    f.append(Paragraph("EDUCATION", s["h"]))
    for d, i, y in r["edu"]:
        f.append(Paragraph(f"{esc(d)}, {esc(i)} ({y})", s["body"]))
    f.append(Paragraph("TECHNICAL SKILLS", s["h"]))
    for g, items in r["skills"].items():
        f.append(Paragraph(f"<b>{esc(g)}:</b> {esc(', '.join(items))}", s["body"]))
    if r["certs"]:
        f.append(Paragraph("CERTIFICATIONS", s["h"]))
        f += [Paragraph(esc(c), s["bullet"], bulletText="•") for c in r["certs"]]
    doc.build(f)


def build_compact_pdf(r, path):
    s = styles()
    doc = SimpleDocTemplate(str(path), pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=12 * mm)
    f = [Paragraph(esc(r["name"].upper()), s["name_left"]), Paragraph(esc(r["title"]), s["title_left"]), Paragraph(esc(_contact(r)), s["contact_left"]),
         Paragraph("Profile", s["h"]), Paragraph(esc(r["summary"]), s["body"]), Paragraph("Core Competencies", s["h"])]
    for g, items in r["skills"].items():
        f.append(Paragraph(f"<b>{esc(g)}</b>: {esc(' • '.join(items))}", s["body"]))
    f.append(Paragraph("Professional Experience", s["h"]))
    for t, c, a, b, bl in r["jobs"]:
        f.append(Paragraph(f"{esc(t)} - {esc(c)}", s["role"]))
        f.append(Paragraph(f"{a} – {b}", s["contact_left"]))
        f += [Paragraph(esc(x), s["bullet"], bulletText="-") for x in bl]
    f.append(Paragraph("Academic Background", s["h"]))
    for d, i, y in r["edu"]:
        f.append(Paragraph(f"{esc(d)}", s["role"]))
        f.append(Paragraph(f"{esc(i)}, {y}", s["body"]))
    if r["certs"]:
        f.append(Paragraph("Licenses & Certifications", s["h"]))
        f += [Paragraph(esc(c), s["bullet"], bulletText="-") for c in r["certs"]]
    doc.build(f)


def build_sidebar_pdf(r, path):
    s = styles()
    W, H = A4
    doc = BaseDocTemplate(str(path), pagesize=A4, leftMargin=0, rightMargin=0, topMargin=0, bottomMargin=0)
    head = Frame(15 * mm, H - 44 * mm, W - 30 * mm, 38 * mm, id="head", leftPadding=0, rightPadding=0)
    left = Frame(15 * mm, 12 * mm, 60 * mm, H - 62 * mm, id="left", leftPadding=0, rightPadding=4)
    right = Frame(82 * mm, 12 * mm, W - 97 * mm, H - 62 * mm, id="right", leftPadding=4, rightPadding=0)
    doc.addPageTemplates([PageTemplate(id="p", frames=[head, left, right])])
    f = [Paragraph(esc(r["name"]), s["name"]), Paragraph(esc(r["title"]), s["title"]), Spacer(1, 3), Paragraph(esc(_contact(r)), s["contact"]), FrameBreak()]
    f.append(Paragraph("SKILLS", s["h"]))
    for g, items in r["skills"].items():
        f.append(Paragraph(f"<b>{esc(g)}</b>", s["small"]))
        f.append(Paragraph(esc(", ".join(items)), s["small"]))
        f.append(Spacer(1, 3))
    f.append(Paragraph("EDUCATION", s["h"]))
    for d, i, y in r["edu"]:
        f += [Paragraph(f"<b>{esc(d)}</b>", s["small"]), Paragraph(esc(i), s["small"]), Paragraph(y, s["small"]), Spacer(1, 3)]
    if r["certs"]:
        f.append(Paragraph("CERTIFICATIONS", s["h"]))
        f += [Paragraph(esc(c), s["small"]) for c in r["certs"]]
    f.append(FrameBreak())
    f += [Paragraph("PROFILE", s["h"]), Paragraph(esc(r["summary"]), s["body"]), Paragraph("WORK EXPERIENCE", s["h"])]
    for t, c, a, b, bl in r["jobs"]:
        f.append(Paragraph(esc(t), s["role"]))
        f.append(Paragraph(f"{esc(c)}  |  {a} - {b}", s["contact_left"]))
        f += [Paragraph(esc(x), s["bullet"], bulletText="•") for x in bl]
    doc.build(f)


def _run(p, text, bold=False, size=10, color=None):
    run = p.add_run(text)
    run.bold, run.font.size = bold, Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def _docx_heading(container, text):
    p = container.add_paragraph()
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(2)
    _run(p, text, True, 11, "0F6E73")


def _docx_body(container, r, with_header=True):
    if with_header:
        p = container.add_paragraph()
        _run(p, r["name"], True, 22, "1B2733")
        p = container.add_paragraph()
        _run(p, r["title"], False, 12, "0F6E73")
        p = container.add_paragraph()
        _run(p, _contact(r), False, 9, "5B6B75")


def build_docx_classic(r, path):
    d = Document()
    _docx_body(d, r)
    _docx_heading(d, "Summary")
    d.add_paragraph(r["summary"])
    _docx_heading(d, "Skills")
    for g, items in r["skills"].items():
        p = d.add_paragraph()
        _run(p, f"{g}: ", True)
        _run(p, ", ".join(items))
    _docx_heading(d, "Work Experience")
    for t, c, a, b, bl in r["jobs"]:
        p = d.add_paragraph()
        _run(p, f"{t}, {c}", True)
        _run(p, f"    {a} - {b}", False, 9, "5B6B75")
        for x in bl:
            d.add_paragraph(x, style="List Bullet")
    _docx_heading(d, "Education")
    for dg, i, y in r["edu"]:
        d.add_paragraph(f"{dg}, {i} ({y})")
    if r["certs"]:
        _docx_heading(d, "Certifications")
        for c in r["certs"]:
            d.add_paragraph(c, style="List Bullet")
    d.save(path)


def build_docx_table(r, path):
    d = Document()
    hp = d.sections[0].header.paragraphs[0]
    _run(hp, f"{r['name']} - Curriculum Vitae", False, 8, "5B6B75")
    _docx_body(d, r)
    t = d.add_table(rows=1, cols=2)
    left, right = t.rows[0].cells
    left.width, right.width = Pt(170), Pt(340)
    left.paragraphs[0].text = ""
    _docx_heading(left, "Skills")
    for g, items in r["skills"].items():
        p = left.add_paragraph()
        _run(p, f"{g}: ", True, 9)
        _run(p, ", ".join(items), False, 9)
    _docx_heading(left, "Education")
    for dg, i, y in r["edu"]:
        p = left.add_paragraph()
        _run(p, dg, True, 9)
        _run(p, f"\n{i}\n{y}", False, 9)
    if r["certs"]:
        _docx_heading(left, "Certifications")
        for c in r["certs"]:
            _run(left.add_paragraph(), c, False, 9)
    right.paragraphs[0].text = ""
    _docx_heading(right, "Profile")
    _run(right.add_paragraph(), r["summary"], False, 9.5)
    _docx_heading(right, "Experience")
    for tt, c, a, b, bl in r["jobs"]:
        p = right.add_paragraph()
        _run(p, tt, True, 10)
        p = right.add_paragraph()
        _run(p, f"{c}  |  {a} - {b}", False, 9, "5B6B75")
        for x in bl:
            right.add_paragraph(x, style="List Bullet")
    d.save(path)


def build_scanned_pdf(r, path):
    import tempfile

    import pypdfium2 as pdfium

    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "src.pdf"
        build_classic_pdf(r, src)
        page = pdfium.PdfDocument(str(src))[0]
        img = page.render(scale=2.2).to_pil().convert("L")
        img.save(path, "PDF", resolution=158.0)


BUILDERS = {"scanned_pdf": build_scanned_pdf, "classic_pdf": build_classic_pdf, "compact_pdf": build_compact_pdf, "sidebar_pdf": build_sidebar_pdf,
            "docx_classic": build_docx_classic, "docx_table": build_docx_table}


def main():
    (OUT / "resumes").mkdir(parents=True, exist_ok=True)
    (OUT / "job_descriptions").mkdir(parents=True, exist_ok=True)
    for r in R:
        BUILDERS[r["layout"]](r, OUT / "resumes" / r["file"])
    for fn, body in JOBS.items():
        (OUT / "job_descriptions" / fn).write_text(body, encoding="utf-8")
    with open(OUT / "resume_manifest.csv", "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["file", "layout", "name", "title", "email"])
        for r in R:
            w.writerow([r["file"], r["layout"], r["name"], r["title"], r["email"]])
    print(f"Generated {len(R)} resumes and {len(JOBS)} job descriptions in {OUT}")


if __name__ == "__main__":
    main()

