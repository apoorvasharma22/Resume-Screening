from __future__ import annotations

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

COLUMNS = [
    ("Rank", "rank"), ("Name", "name"), ("Email", "email"), ("Phone", "phone"), ("Match score (%)", "score"),
    ("Verdict", "verdict"), ("Status", "status"), ("Experience (yrs)", "experience_years"),
    ("Highest education", "highest_education"), ("Matched required skills", "matched_required"),
    ("Missing required skills", "missing_required"), ("Missing preferred skills", "missing_preferred"),
    ("Certifications", "certifications"), ("All skills", "skills"), ("Summary", "summary"),
]


def _cell(v):
    return ", ".join(v) if isinstance(v, list) else ("" if v is None else v)


def flatten(item: dict) -> dict:
    c, a = item["candidate"], item["skill_analysis"]
    return {
        "rank": item["rank"], "name": c["name"], "email": c["email"], "phone": c["phone"], "score": item["score"],
        "verdict": item["verdict"], "status": item["status"], "experience_years": c["experience_years"],
        "highest_education": c["highest_education"], "matched_required": a["matched_required"],
        "missing_required": a["missing_required"], "missing_preferred": a["missing_preferred"],
        "certifications": c["certifications"], "skills": c["skills"], "summary": item["summary"],
    }


def to_csv(items: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([h for h, _ in COLUMNS])
    for it in items:
        row = flatten(it)
        w.writerow([_cell(row[k]) for _, k in COLUMNS])
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def to_xlsx(items: list[dict], job: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Shortlist"
    ws["A1"] = f"{job['title']}" + (f" - {job['company']}" if job.get("company") else "")
    ws["A1"].font = Font(size=14, bold=True, color="0B2A33")
    ws["A2"] = f"{len(items)} candidates, ranked by match score"
    ws["A2"].font = Font(italic=True, color="5B6B75")

    header_row = 4
    fill = PatternFill("solid", fgColor="0B2A33")
    thin = Side(style="thin", color="D8DEE3")
    for i, (h, _) in enumerate(COLUMNS, 1):
        c = ws.cell(row=header_row, column=i, value=h)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[header_row].height = 30

    for r, it in enumerate(items, header_row + 1):
        row = flatten(it)
        for i, (_, k) in enumerate(COLUMNS, 1):
            c = ws.cell(row=r, column=i, value=_cell(row[k]))
            c.alignment = Alignment(vertical="top", wrap_text=k in ("summary", "all skills", "skills", "certifications", "matched_required", "missing_required"))
            c.border = Border(bottom=thin)
        score = row["score"]
        colour = "CFEFE3" if score >= 75 else "E4F3D3" if score >= 60 else "FFF0C7" if score >= 40 else "FBD9D0"
        ws.cell(row=r, column=5).fill = PatternFill("solid", fgColor=colour)
        ws.cell(row=r, column=5).font = Font(bold=True)

    widths = [6, 24, 30, 16, 13, 15, 12, 12, 14, 38, 34, 34, 34, 50, 60]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=header_row + 1, column=3)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(COLUMNS))}{header_row + max(len(items), 1)}"

    info = wb.create_sheet("Job requirements")
    rows = [("Title", job["title"]), ("Company", job.get("company") or ""),
            ("Required skills", ", ".join(job["required_skills"])), ("Preferred skills", ", ".join(job["preferred_skills"])),
            ("Minimum experience (yrs)", job.get("min_experience") or ""), ("Education level", job.get("education_level") or "")]
    for i, (k, v) in enumerate(rows, 1):
        info.cell(row=i, column=1, value=k).font = Font(bold=True)
        info.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True, vertical="top")
    info.column_dimensions["A"].width = 28
    info.column_dimensions["B"].width = 90

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

