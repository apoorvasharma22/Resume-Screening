from __future__ import annotations

import hashlib
import io
import re
import unicodedata
from dataclasses import dataclass, field

from app.config import get_settings
from app.exceptions import FileTooLargeError, ResumeParseError, UnsupportedFileTypeError
from app.logging_config import get_logger

log = get_logger(__name__)

ALLOWED_EXTENSIONS = {".pdf", ".docx"}
MIN_TEXT_CHARS = 120


@dataclass
class ParsedDocument:
    text: str
    file_type: str
    method: str
    file_hash: str
    warnings: list[str] = field(default_factory=list)


def validate_upload(filename: str, content: bytes) -> str:
    settings = get_settings()
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED_EXTENSIONS:
        raise UnsupportedFileTypeError(f"'{filename}': only PDF and DOCX resumes are supported (got '{ext or 'no extension'}').")
    if len(content) == 0:
        raise ResumeParseError(f"'{filename}' is empty.")
    if len(content) > settings.max_upload_mb * 1024 * 1024:
        raise FileTooLargeError(f"'{filename}' is larger than {settings.max_upload_mb} MB.")
    if ext == ".pdf" and not content.lstrip()[:5].startswith(b"%PDF"):
        raise ResumeParseError(f"'{filename}' has a .pdf extension but is not a valid PDF file.")
    if ext == ".docx" and not content.startswith(b"PK"):
        raise ResumeParseError(f"'{filename}' has a .docx extension but is not a valid Word document.")
    return ext.lstrip(".")


def sha256(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


_BULLETS = "•●▪■◦○◆◇►▶➤✔✓‣∙·"


def clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"\(cid:\d+\)", "- ", text)
    text = text.replace("\x00", "").replace("\u00a0", " ")
    for b in _BULLETS:
        text = text.replace(b, "- ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return "\n".join(line.strip() for line in text.splitlines()).strip()


def _corridor(words: list[dict], page_width: float) -> float | None:
    total = len(words)
    if total < 30:
        return None
    steps = 240
    valid: list[int] = []
    for i in range(int(steps * 0.15), int(steps * 0.85)):
        x = page_width * i / steps
        if any(w["x0"] < x < w["x1"] for w in words):
            continue
        left = sum(1 for w in words if w["x1"] <= x)
        if min(left, total - left) >= 0.12 * total:
            valid.append(i)
    if not valid:
        return None
    runs, run = [], [valid[0]]
    for i in valid[1:]:
        if i == run[-1] + 1:
            run.append(i)
        else:
            runs.append(run)
            run = [i]
    runs.append(run)
    best = max(runs, key=len)
    if len(best) < 3:
        return None
    return page_width * (best[0] + best[-1]) / 2 / steps


def _pdf_page_text(page) -> tuple[str, bool]:
    words = page.extract_words(x_tolerance=2, y_tolerance=3)
    height, width = float(page.height), float(page.width)
    gutter = None
    for cutoff in (0.30, 0.22, 0.15, 0.08, 0.0):
        gutter = _corridor([w for w in words if w["top"] >= cutoff * height], width)
        if gutter is not None:
            break
    if gutter is None:
        return page.extract_text() or "", False

    band_bottom = max((w["bottom"] for w in words if w["x0"] < gutter < w["x1"]), default=0.0)
    parts = []
    if band_bottom:
        parts.append(page.crop((0, 0, width, band_bottom + 1)).extract_text() or "")
    top = band_bottom + 1 if band_bottom else 0
    for box in ((0, top, gutter, height), (gutter, top, width, height)):
        parts.append(page.crop(box).extract_text() or "")
    return "\n".join(p for p in parts if p.strip()), True


def _ocr_pdf(content: bytes) -> str:
    try:
        import pypdfium2 as pdfium
        import pytesseract
    except ImportError as exc:
        raise ResumeParseError("This PDF looks scanned and OCR libraries are not installed.") from exc
    try:
        pdf = pdfium.PdfDocument(content)
        out = []
        for i in range(min(len(pdf), 6)):
            image = pdf[i].render(scale=2.5).to_pil()
            out.append(pytesseract.image_to_string(image))
        return "\n".join(out)
    except pytesseract.TesseractNotFoundError as exc:
        raise ResumeParseError(
            "This PDF looks scanned but the Tesseract OCR engine is not installed on the server "
            "(it is included in the Docker image)."
        ) from exc


def parse_pdf(content: bytes, filename: str) -> tuple[str, str, list[str]]:
    import pdfplumber
    from pdfminer.pdfdocument import PDFPasswordIncorrect

    warnings: list[str] = []
    try:
        with pdfplumber.open(io.BytesIO(content)) as pdf:
            if not pdf.pages:
                raise ResumeParseError(f"'{filename}' contains no pages.")
            pages, used_columns = [], False
            for page in pdf.pages:
                text, cols = _pdf_page_text(page)
                pages.append(text)
                used_columns = used_columns or cols
    except ResumeParseError:
        raise
    except Exception as exc:
        log.warning("PDF open failed for %s: %s", filename, exc)
        raise ResumeParseError(f"'{filename}' could not be read (corrupted or password-protected PDF).") from exc

    text = clean_text("\n".join(pages))
    method = "pdf-columns" if used_columns else "pdf-text"

    if len(text) < MIN_TEXT_CHARS:
        if not get_settings().ocr_enabled:
            raise ResumeParseError(f"'{filename}' has no selectable text (scanned PDF) and OCR is disabled.")
        log.info("Falling back to OCR for %s (only %d chars of text)", filename, len(text))
        text = clean_text(_ocr_pdf(content))
        method = "pdf-ocr"
        warnings.append("Scanned PDF: text was recovered with OCR and may contain errors.")
    return text, method, warnings


def _iter_docx_blocks(doc):
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in doc.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, doc)
        elif child.tag.endswith("}tbl"):
            yield Table(child, doc)


def _table_lines(table) -> list[str]:
    lines: list[str] = []
    for row in table.rows:
        seen: set[int] = set()
        for cell in row.cells:
            if id(cell._tc) in seen:
                continue
            seen.add(id(cell._tc))
            for p in cell.paragraphs:
                if p.text.strip():
                    lines.append(p.text)
            for nested in cell.tables:
                lines.extend(_table_lines(nested))
    return lines


def parse_docx(content: bytes, filename: str) -> tuple[str, str, list[str]]:
    from docx import Document
    from docx.table import Table

    try:
        doc = Document(io.BytesIO(content))
    except Exception as exc:
        log.warning("DOCX open failed for %s: %s", filename, exc)
        raise ResumeParseError(f"'{filename}' could not be read (corrupted Word file).") from exc

    lines: list[str] = []
    for section in doc.sections:
        lines.extend(p.text for p in section.header.paragraphs if p.text.strip())
    for block in _iter_docx_blocks(doc):
        if isinstance(block, Table):
            lines.extend(_table_lines(block))
        elif block.text.strip():
            lines.append(block.text)

    text = clean_text("\n".join(lines))
    if len(text) < 40:
        raise ResumeParseError(f"'{filename}' contains almost no text.")
    return text, "docx", []


def parse_resume(filename: str, content: bytes) -> ParsedDocument:
    ext = validate_upload(filename, content)
    if ext == "pdf":
        text, method, warnings = parse_pdf(content, filename)
    else:
        text, method, warnings = parse_docx(content, filename)
    if len(text) < 40:
        raise ResumeParseError(f"'{filename}' contains almost no readable text.")
    log.info("Parsed %s (%s, %d chars)", filename, method, len(text))
    return ParsedDocument(text=text, file_type=ext, method=method, file_hash=sha256(content), warnings=warnings)

